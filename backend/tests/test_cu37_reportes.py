"""CU-37 · Generar reportes de gestión.

Realiza el **RF36** —«exportar los reportes en formato PDF y Excel»—, que
hasta el 19/09/2026 no lo cubría nadie: CU-36 muestra los indicadores en
pantalla y ahí terminaba.

Lo que más importa cubrir
-------------------------
- **Que las seis consultas CORRAN.** Es lo primero y lo más valioso: son seis
  consultas con entre tres y siete JOIN, y un nombre de columna equivocado no
  se nota hasta que alguien pide ese reporte. Escribiéndolas ya aparecieron
  dos: `Reserva` no tiene `codigo`, y `DetalleVenta` no guarda `subtotal`.
- **Que los archivos sean archivos.** Un PDF que empieza con otra cosa que
  `%PDF` y un XLSX que no es un ZIP son bytes que el navegador descarga y
  ninguna aplicación abre.
- **Que el encargado no vea la red entera.** Se le fuerza la sucursal en vez
  de confiar en el parámetro: quitar `?sucursal_id=` de la URL no puede darle
  las ventas de las demás.
- **Que un reporte vacío no sea un error.** En un período sin movimiento no
  pasó nada, y eso es una respuesta válida.
"""

from __future__ import annotations

from datetime import date, timedelta

import pytest
from fastapi.testclient import TestClient

REPORTES = "/api/v1/reportes"
CATALOGO = f"{REPORTES}/catalogo"
EMPLEADOS = "/api/v1/organizacion/empleados"
SUCURSALES = "/api/v1/organizacion/sucursales"

#: Los seis del RF36. Si alguien agrega uno, esta lista lo obliga a probarlo.
TIPOS = ("ventas", "inventario", "movimientos", "reservas", "rendimiento", "compras")


@pytest.fixture
def sucursal(api: TestClient, cabeceras_admin: dict[str, str]) -> int:
    r = api.post(
        SUCURSALES,
        headers=cabeceras_admin,
        json={
            "ciudad_id": 1,
            "nombre": "Centro",
            "direccion": "Avenida Centro 100",
            "telefono": None,
            "horario_apertura": "09:00",
            "horario_cierre": "20:00",
            "capacidad_vestidores": 2,
            "activa": True,
        },
    )
    assert r.status_code == 201, r.text
    return r.json()["id"]


@pytest.fixture
def encargado(api: TestClient, cabeceras_admin: dict[str, str], sucursal: int) -> dict:
    correo = "encargado.reportes@violetboutique.bo"
    clave = "Encargado12"
    r = api.post(
        EMPLEADOS,
        headers=cabeceras_admin,
        json={
            "nombres": "Luz",
            "apellidos": "Vargas",
            "correo": correo,
            "contrasena": clave,
            "documento": "5544339",
            "telefono": "70000000",
            "cargo": "ENCARGADO",
            "sucursal_id": sucursal,
            "fecha_ingreso": (date.today() - timedelta(days=30)).isoformat(),
        },
    )
    assert r.status_code == 201, r.text
    entrada = api.post("/api/v1/auth/login", json={"correo": correo, "contrasena": clave})
    return {"Authorization": f"Bearer {entrada.json()['access_token']}"}


# --- Que las consultas corran ----------------------------------------------


@pytest.mark.parametrize("tipo", TIPOS)
def test_cada_reporte_se_genera_en_PDF(
    api: TestClient, cabeceras_admin: dict[str, str], tipo: str
) -> None:
    """Lo más valioso del archivo: seis consultas con muchos JOIN.

    Un nombre de columna equivocado no se nota hasta que alguien pide ese
    reporte, y para entonces es delante de un tribunal.
    """
    r = api.get(f"{REPORTES}/{tipo}.pdf", headers=cabeceras_admin)
    assert r.status_code == 200, r.text
    assert r.content.startswith(b"%PDF"), "no es un PDF"
    assert r.headers["content-type"] == "application/pdf"
    assert f"{tipo}-" in r.headers["content-disposition"]


@pytest.mark.parametrize("tipo", TIPOS)
def test_cada_reporte_se_genera_en_EXCEL(
    api: TestClient, cabeceras_admin: dict[str, str], tipo: str
) -> None:
    r = api.get(f"{REPORTES}/{tipo}.xlsx", headers=cabeceras_admin)
    assert r.status_code == 200, r.text
    # Un .xlsx es un ZIP. Si no empieza con PK, ninguna aplicación lo abre.
    assert r.content.startswith(b"PK"), "no es un XLSX"
    assert "spreadsheetml" in r.headers["content-type"]


def test_un_reporte_VACIO_no_es_un_error(
    api: TestClient, cabeceras_admin: dict[str, str]
) -> None:
    """En un período sin movimiento no pasó nada, y eso es una respuesta."""
    r = api.get(
        f"{REPORTES}/ventas.pdf",
        headers=cabeceras_admin,
        params={"desde": "2020-01-01", "hasta": "2020-01-31"},
    )
    assert r.status_code == 200
    assert r.content.startswith(b"%PDF")


def test_el_excel_se_puede_abrir_y_tiene_los_encabezados(
    api: TestClient, cabeceras_admin: dict[str, str]
) -> None:
    """Que empiece con PK no alcanza: se abre y se mira adentro."""
    from io import BytesIO

    from openpyxl import load_workbook

    r = api.get(f"{REPORTES}/inventario.xlsx", headers=cabeceras_admin)
    libro = load_workbook(BytesIO(r.content))
    hoja = libro.active
    textos = [c.value for fila in hoja.iter_rows(max_row=8) for c in fila if c.value]
    assert "Reporte de inventario" in textos
    assert "Disponible" in textos
    # El alcance SIEMPRE se imprime: un reporte sin decir de qué sucursal es
    # un papel que no se puede archivar.
    assert any("Sucursal: todas" in str(t) for t in textos)


# --- El periodo -------------------------------------------------------------


def test_el_ultimo_dia_del_periodo_entra_ENTERO(
    api: TestClient, cabeceras_admin: dict[str, str]
) -> None:
    """El defecto clásico de los reportes por fecha.

    Comparar `<= hasta` deja fuera todo lo que pasó ese día después de
    medianoche. Se convierte a «< el día siguiente». Acá se comprueba que el
    subtítulo diga el día que el usuario pidió, no el siguiente.
    """
    from io import BytesIO

    from openpyxl import load_workbook

    hoy = date.today()
    r = api.get(
        f"{REPORTES}/ventas.xlsx",
        headers=cabeceras_admin,
        params={"desde": hoy.isoformat(), "hasta": hoy.isoformat()},
    )
    hoja = load_workbook(BytesIO(r.content)).active
    textos = " ".join(
        str(c.value) for fila in hoja.iter_rows(max_row=5) for c in fila if c.value
    )
    esperado = hoy.strftime("%d/%m/%Y")
    assert f"{esperado} al {esperado}" in textos


def test_una_fecha_inicial_posterior_a_la_final_se_rechaza(
    api: TestClient, cabeceras_admin: dict[str, str]
) -> None:
    r = api.get(
        f"{REPORTES}/ventas.pdf",
        headers=cabeceras_admin,
        params={"desde": "2026-09-30", "hasta": "2026-09-01"},
    )
    assert r.status_code == 422


# --- El catálogo ------------------------------------------------------------


def test_el_catalogo_lista_los_seis(
    api: TestClient, cabeceras_admin: dict[str, str]
) -> None:
    """Existe para que la pantalla no tenga la lista escrita a mano."""
    r = api.get(CATALOGO, headers=cabeceras_admin)
    assert r.status_code == 200
    tipos = {f["tipo"] for f in r.json()}
    assert tipos == set(TIPOS)
    inventario = next(f for f in r.json() if f["tipo"] == "inventario")
    # El inventario es una foto de ahora: no usa período, y la pantalla tiene
    # que saberlo para esconder el selector de fechas.
    assert inventario["usa_periodo"] is False


def test_un_reporte_que_no_existe_da_404(
    api: TestClient, cabeceras_admin: dict[str, str]
) -> None:
    r = api.get(f"{REPORTES}/inventado.pdf", headers=cabeceras_admin)
    assert r.status_code == 404
    assert "disponibles" in r.json()["detail"].lower()


# --- Quién puede, y sobre qué ----------------------------------------------


def test_EL_ENCARGADO_SOLO_VE_SU_SUCURSAL(
    api: TestClient, encargado: dict, sucursal: int
) -> None:
    """Se le FUERZA el filtro en vez de confiar en el parámetro.

    Sin esto, quitar `?sucursal_id=` de la URL le daría las ventas de toda la
    red — que es exactamente el dato que no corresponde que vea.
    """
    from io import BytesIO

    from openpyxl import load_workbook

    r = api.get(f"{REPORTES}/ventas.xlsx", headers=encargado)
    assert r.status_code == 200
    hoja = load_workbook(BytesIO(r.content)).active
    textos = " ".join(
        str(c.value) for fila in hoja.iter_rows(max_row=5) for c in fila if c.value
    )
    assert "Sucursal: Centro" in textos
    assert "Sucursal: todas" not in textos


def test_el_encargado_NO_puede_pedir_otra_sucursal(
    api: TestClient, encargado: dict, cabeceras_admin: dict, sucursal: int
) -> None:
    """Aunque mande el parámetro a mano, se le impone la suya."""
    from io import BytesIO

    from openpyxl import load_workbook

    otra = api.post(
        SUCURSALES,
        headers=cabeceras_admin,
        json={
            "ciudad_id": 1,
            "nombre": "Norte",
            "direccion": "Avenida Norte 200",
            "telefono": None,
            "horario_apertura": "09:00",
            "horario_cierre": "20:00",
            "capacidad_vestidores": 1,
            "activa": True,
        },
    ).json()["id"]

    r = api.get(
        f"{REPORTES}/ventas.xlsx", headers=encargado, params={"sucursal_id": otra}
    )
    hoja = load_workbook(BytesIO(r.content)).active
    textos = " ".join(
        str(c.value) for fila in hoja.iter_rows(max_row=5) for c in fila if c.value
    )
    assert "Sucursal: Centro" in textos
    assert "Sucursal: Norte" not in textos


def test_un_cliente_no_descarga_reportes(
    api: TestClient, cabeceras_cliente: dict[str, str]
) -> None:
    assert api.get(f"{REPORTES}/ventas.pdf", headers=cabeceras_cliente).status_code == 403


def test_sin_token_no_hay_reportes(api: TestClient) -> None:
    assert api.get(f"{REPORTES}/ventas.pdf").status_code == 401
