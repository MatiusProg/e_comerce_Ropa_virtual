# Genera "SI2 Ex2 Grupo 15.docx" (ya NO se usa: el Word se edita a mano; ver editar_22_10.py) a partir de la plantilla del 1er parcial (Defensa.docx):
# conserva portada, índice, estilos, encabezados y pie; reemplaza el cuerpo.
import copy, sys
import docx
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from contenido import CONTENIDO

PLANTILLA = sys.argv[1]
SALIDA = sys.argv[2]

doc = docx.Document(PLANTILLA)
body = doc.element.body
hijos = list(body.iterchildren())
# 0 = portada, 1 = índice (sdt); se borra todo lo demás salvo el sectPr final
for el in hijos[2:]:
    if el.tag != qn('w:sectPr'):
        body.remove(el)

# ---------- portada: título nuevo ----------
NUEVO_TITULO = ("Violet Boutique: plataforma SaaS multitenant de comercio electrónico para empresas "
                "de venta de ropa, con operación offline en web y móvil, sincronización de datos y "
                "reportes generativos con inteligencia artificial, desarrollada con la metodología "
                "ICONIX y modelos UM")
portada = hijos[0]
ts = portada.iter(qn('w:t'))
ts = list(ts)
for i, t in enumerate(ts):
    if t.text and t.text.startswith('lataforma inteligente'):
        t.text = NUEVO_TITULO[1:]
        if i > 0 and ts[i - 1].text == 'P':
            ts[i - 1].text = NUEVO_TITULO[0]

# ---------- listas numeradas: una <w:num> nueva por lista para que reinicie ----------
numbering = doc.part.numbering_part.element
ABS_DECIMAL = '6'      # abstractNum decimal '%1.' de la plantilla
NUM_VIÑETA = '25'      # viñeta usada en todo el documento del 1er parcial
_sig_num = [300]


def nueva_lista_numerada():
    _sig_num[0] += 1
    n = OxmlElement('w:num')
    n.set(qn('w:numId'), str(_sig_num[0]))
    a = OxmlElement('w:abstractNumId')
    a.set(qn('w:val'), ABS_DECIMAL)
    n.append(a)
    o = OxmlElement('w:lvlOverride')
    o.set(qn('w:ilvl'), '0')
    s = OxmlElement('w:startOverride')
    s.set(qn('w:val'), '1')
    o.append(s)
    n.append(o)
    numbering.append(n)
    return str(_sig_num[0])


# ---------- piezas ----------
def _formato_cuerpo(p, justificar=True, sangria_cero=False):
    pf = p.paragraph_format
    pf.line_spacing = 1.15
    if justificar:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if sangria_cero:
        pf.first_line_indent = 0


def _runs(p, texto, tam=None, color=None):
    """**negrita** dentro del texto."""
    partes = texto.split('**')
    for i, parte in enumerate(partes):
        if not parte:
            continue
        r = p.add_run(parte)
        r.bold = (i % 2 == 1) or None
        if tam:
            r.font.size = Pt(tam)
        if color:
            r.font.color.rgb = color
    return p


def salto_de_pagina():
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def titulo(estilo, texto):
    doc.add_paragraph(texto, style=estilo)


def parrafo(texto):
    p = doc.add_paragraph()
    _formato_cuerpo(p)
    _runs(p, texto)


def lista(items, numerada=False):
    num = nueva_lista_numerada() if numerada else NUM_VIÑETA
    for it in items:
        p = doc.add_paragraph(style='List Paragraph')
        pPr = p._p.get_or_add_pPr()
        numPr = OxmlElement('w:numPr')
        il = OxmlElement('w:ilvl'); il.set(qn('w:val'), '0')
        ni = OxmlElement('w:numId'); ni.set(qn('w:val'), num)
        numPr.append(il); numPr.append(ni)
        pPr.insert(1, numPr)
        _formato_cuerpo(p)
        _runs(p, it)


def pendiente(texto):
    p = doc.add_paragraph()
    _formato_cuerpo(p)
    r = p.add_run('[Pendiente] ' + texto)
    r.italic = True
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW


BORDE = 'D1D9E0'
ROJO = RGBColor(0xEE, 0x00, 0x00)


def _celda(celda, texto, ancho, encabezado, cebra, negrita):
    tcPr = celda._tc.get_or_add_tcPr()
    for hijo in list(tcPr):
        tcPr.remove(hijo)
    w = OxmlElement('w:tcW'); w.set(qn('w:w'), str(ancho)); w.set(qn('w:type'), 'dxa'); tcPr.append(w)
    b = OxmlElement('w:tcBorders')
    for lado in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{lado}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '6'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), BORDE)
        b.append(e)
    tcPr.append(b)
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto')
    s.set(qn('w:fill'), 'F6F8FA' if cebra else 'FFFFFF'); tcPr.append(s)
    m = OxmlElement('w:tcMar')
    for lado, v in (('top', 90), ('left', 195), ('bottom', 90), ('right', 195)):
        e = OxmlElement(f'w:{lado}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tcPr.append(m)
    va = OxmlElement('w:vAlign'); va.set(qn('w:val'), 'center'); tcPr.append(va)
    p = celda.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.first_line_indent = 0
    if encabezado:
        r = p.add_run(texto); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = ROJO
    else:
        _runs(p, f'**{texto}**' if negrita and '**' not in texto else texto, tam=11)


def tabla(encabezados, filas, anchos):
    total = 9782
    assert sum(anchos) == total, (encabezados, sum(anchos))
    t = doc.add_table(rows=1 + len(filas), cols=len(encabezados))
    tblPr = t._tbl.tblPr
    for tag in ('w:tblStyle', 'w:tblW', 'w:tblInd', 'w:tblLook'):
        for e in tblPr.findall(qn(tag)):
            tblPr.remove(e)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(total)); tw.set(qn('w:type'), 'dxa'); tblPr.append(tw)
    ti = OxmlElement('w:tblInd'); ti.set(qn('w:w'), '-434'); ti.set(qn('w:type'), 'dxa'); tblPr.append(ti)
    grid = t._tbl.tblGrid
    for gc, a in zip(grid.findall(qn('w:gridCol')), anchos):
        gc.set(qn('w:w'), str(a))
    # encabezado repetido en cada página
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader'); trPr.append(th)
    for j, h in enumerate(encabezados):
        _celda(t.rows[0].cells[j], h, anchos[j], True, False, False)
    for i, fila in enumerate(filas, start=1):
        for j, v in enumerate(fila):
            _celda(t.rows[i].cells[j], v, anchos[j], False, i % 2 == 0, j == 0)
    doc.add_paragraph()


ACCIONES = {
    'salto': lambda a: salto_de_pagina(),
    'cap': lambda a: titulo('CAPITULO', a[1]),
    'h1': lambda a: titulo('Titulo1', a[1]),
    'h2': lambda a: titulo('Titulo2', a[1]),
    'h3': lambda a: titulo('Titulo3', a[1]),
    'p': lambda a: parrafo(a[1]),
    'ul': lambda a: lista(a[1]),
    'ol': lambda a: lista(a[1], numerada=True),
    'pend': lambda a: pendiente(a[1]),
    'tabla': lambda a: tabla(a[1], a[2], a[3]),
}

for accion in CONTENIDO:
    ACCIONES[accion[0]](accion)

# el sectPr tiene que quedar último
sect = body.find(qn('w:sectPr'))
body.remove(sect); body.append(sect)

# Word actualiza el índice al abrir
settings = doc.settings.element
if settings.find(qn('w:updateFields')) is None:
    u = OxmlElement('w:updateFields'); u.set(qn('w:val'), 'true'); settings.append(u)

doc.save(SALIDA)
print('ok', SALIDA)
