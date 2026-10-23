# Correcciones de la ing del 22/10 aplicadas sobre el Word ya editado a mano (no se regenera con gen_doc.py,
# que pisaría los cambios de Luis):
#   - sin ciclos ni priorización: un solo documento al que cada presentación le agrega casos de uso;
#   - los casos de uso se numeran desde el CU-01 (CU-43..CU-77 -> CU-01..CU-35);
#   - los del 1er parcial se citan como «del primer parcial» y se listan en el Anexo 1;
#   - la lista de casos de uso lleva la columna Plataforma (web, móvil o ambos).
# Uso: python editar_22_10.py <documento.docx> <plantilla_1er_parcial.docx>
import copy, io, re, sys
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

DOC, P1 = sys.argv[1], sys.argv[2]
doc = docx.Document(DOC)
body = doc.element.body


def texto(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def buscar(pred, desde=None):
    hijos = list(body.iterchildren())
    i0 = hijos.index(desde) + 1 if desde is not None else 0
    for el in hijos[i0:]:
        if pred(el):
            return el
    raise LookupError('no encontrado')


def parrafo_que_empieza(prefijo, desde=None):
    return buscar(lambda e: e.tag == qn('w:p') and texto(e).startswith(prefijo), desde)


def reemplazar_en_parrafo(p, patron, funcion):
    """Reemplaza sobre el texto completo del párrafo aunque esté partido en varios runs; el formato de cada
    run se conserva (el reemplazo queda en el run donde empieza la coincidencia)."""
    ts = list(p.iter(qn('w:t')))
    if not ts:
        return 0
    completo = ''.join(t.text or '' for t in ts)
    coincidencias = list(re.finditer(patron, completo))
    if not coincidencias:
        return 0
    # mapa: posición en el texto completo -> (índice de w:t, posición dentro de él)
    mapa = []
    for k, t in enumerate(ts):
        for j in range(len(t.text or '')):
            mapa.append((k, j))
    textos = [list(t.text or '') for t in ts]
    for m in reversed(coincidencias):
        nuevo = funcion(m)
        a, b = m.start(), m.end()
        for pos in range(b - 1, a, -1):
            k, j = mapa[pos]
            textos[k][j] = ''
        k, j = mapa[a]
        textos[k][j] = nuevo
    for t, chars in zip(ts, textos):
        t.text = ''.join(chars)
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    return len(coincidencias)


def poner_texto(p, nuevo):
    """Deja el párrafo con un único texto, con **negrita** marcada, conservando el formato del primer run."""
    runs = p.findall(qn('w:r'))
    base = copy.deepcopy(runs[0]) if runs else None
    for r in runs:
        p.remove(r)
    for i, parte in enumerate(nuevo.split('**')):
        if not parte:
            continue
        r = copy.deepcopy(base) if base is not None else OxmlElement('w:r')
        for hijo in list(r):
            if hijo.tag != qn('w:rPr'):
                r.remove(hijo)
        rpr = r.find(qn('w:rPr'))
        if rpr is None:
            rpr = OxmlElement('w:rPr'); r.insert(0, rpr)
        for b in rpr.findall(qn('w:b')):
            rpr.remove(b)
        if i % 2 == 1:
            rpr.insert(0, OxmlElement('w:b'))
        t = OxmlElement('w:t'); t.text = parte
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        r.append(t)
        p.append(r)


# ---------- tablas con el formato del documento (igual que gen_doc.py) ----------
BORDE = 'D1D9E0'
ROJO = RGBColor(0xEE, 0x00, 0x00)


def _runs(p, txt, tam=None):
    for i, parte in enumerate(txt.split('**')):
        if parte:
            r = p.add_run(parte)
            r.bold = (i % 2 == 1) or None
            if tam:
                r.font.size = Pt(tam)


def _celda(celda, txt, ancho, encabezado, cebra, negrita):
    tcPr = celda._tc.get_or_add_tcPr()
    for hijo in list(tcPr):
        tcPr.remove(hijo)
    w = OxmlElement('w:tcW'); w.set(qn('w:w'), str(ancho)); w.set(qn('w:type'), 'dxa'); tcPr.append(w)
    b = OxmlElement('w:tcBorders')
    for lado in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{lado}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '6'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), BORDE)
        b.append(e)
    tcPr.append(b)
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto')
    s.set(qn('w:fill'), 'F6F8FA' if cebra else 'FFFFFF'); tcPr.append(s)
    m = OxmlElement('w:tcMar')
    for lado, v in (('top', 90), ('left', 195), ('bottom', 90), ('right', 195)):
        e = OxmlElement(f'w:{lado}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tcPr.append(m)
    va = OxmlElement('w:vAlign'); va.set(qn('w:val'), 'center'); tcPr.append(va)
    p = celda.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.first_line_indent = 0
    if encabezado:
        r = p.add_run(txt); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = ROJO
    else:
        _runs(p, f'**{txt}**' if negrita and '**' not in txt else txt, tam=11)


def tabla(encabezados, filas, anchos):
    """Crea la tabla al final del documento y devuelve su elemento, para moverlo donde haga falta."""
    total = 9782
    assert sum(anchos) == total, (encabezados, sum(anchos))
    t = doc.add_table(rows=1 + len(filas), cols=len(encabezados))
    tblPr = t._tbl.tblPr
    for tag in ('w:tblStyle', 'w:tblW', 'w:tblInd', 'w:tblLook'):
        for e in tblPr.findall(qn(tag)):
            tblPr.remove(e)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(total)); tw.set(qn('w:type'), 'dxa'); tblPr.append(tw)
    ti = OxmlElement('w:tblInd'); ti.set(qn('w:w'), '-434'); ti.set(qn('w:type'), 'dxa'); tblPr.append(ti)
    for gc, a in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), anchos):
        gc.set(qn('w:w'), str(a))
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))
    for j, h in enumerate(encabezados):
        _celda(t.rows[0].cells[j], h, anchos[j], True, False, False)
    for i, fila in enumerate(filas, start=1):
        for j, v in enumerate(fila):
            _celda(t.rows[i].cells[j], v, anchos[j], False, i % 2 == 0, j == 0)
    return t._tbl


def nuevo_parrafo(txt='', estilo=None):
    p = doc.add_paragraph(style=estilo) if estilo else doc.add_paragraph()
    if txt:
        if estilo:
            p.add_run(txt)
        else:
            p.paragraph_format.line_spacing = 1.15
            p.alignment = 3  # justificado
            _runs(p, txt)
    return p._p


def insertar_despues(ancla, elementos):
    for el in elementos:
        ancla.addnext(el)
        ancla = el
    return ancla


# =====================================================================
# 1) Referencias a casos de uso del primer parcial, antes de renumerar (después chocarían con los nuevos)
p = parrafo_que_empieza('La docente los pidió en clase')
reemplazar_en_parrafo(p, r'\(CU-32\)', lambda m: '(CU-32 del primer parcial)')
p = parrafo_que_empieza('Desde el primer parcial se registran transferencias')
reemplazar_en_parrafo(p, r'\(CU-15\)', lambda m: '(CU-15 del primer parcial)')
for el in body.iter(qn('w:p')):
    reemplazar_en_parrafo(el, r'Evoluciona el (CU-\d\d)\.',
                          lambda m: f'Evoluciona el {m.group(1)} del primer parcial (Anexo 1).')

# =====================================================================
# 2) 1.6: sin ciclos
p = parrafo_que_empieza('Cómo se aplica en este proyecto.')
poner_texto(p, '**Cómo se aplica en este proyecto.** Siguiendo la indicación de la docente, no se priorizan '
               'los casos de uso ni se trabaja por ciclos separados. Hay un solo documento, y en cada '
               'presentación se le agregan casos de uso que recorren todo el proceso: requisitos, robustez, '
               'secuencia y diseño de clases, implementación y pruebas. Los casos de uso se numeran desde el '
               'CU-01; los del primer parcial que no se rehacen se referencian en el Anexo 1. Los diagramas se '
               'terminan sobre el código ya escrito, para que representen lo realmente implementado, y se '
               'modelan en Enterprise Architect.')
vieja = p.getnext()
while vieja.tag != qn('w:tbl'):
    vieja = vieja.getnext()
nueva = tabla(['Presentación', 'Fecha', 'Casos de uso que se agregan'], [
    ['Presentación 1', 'dom 25/10', 'CU-01 a CU-12: empresas, planes y suscripciones; aislamiento con RLS; '
     'contratación desde la web y marca de la tienda; capítulos I y II del documento'],
    ['Presentación 2', 'dom 01/11', 'CU-13 a CU-24: contrato de sincronización; PWA con IndexedDB; Flutter '
     'con SQLite; ventas offline, cambios de prenda, devoluciones web y avisos por correo'],
    ['Presentación 3', 'dom 08/11', 'CU-25 a CU-35: reportes con IA y voz; analítica con ML; traspasos y '
     'probador por talla; exportación de datos'],
    ['Defensa', 'mar 10/11, 7:00', 'Sistema completo desplegado y documento final unificado'],
], [1900, 1700, 6182])
vieja.addprevious(nueva)
body.remove(vieja)

# =====================================================================
# 3) 2.5: una sola lista de casos de uso, sin separar por ciclo, con la columna Plataforma
W, M, WM, BE = 'Web', 'Móvil', 'Web y móvil', 'Backend'
CASOS = [
    ['CU-01', 'Gestionar empresas', 'Administrador de la plataforma', W, '—', 'RF-01, RF-02, RF-03'],
    ['CU-02', 'Gestionar planes', 'Administrador de la plataforma', W, '—', 'RF-04'],
    ['CU-03', 'Gestionar suscripciones', 'Administrador de la plataforma', W, '—', 'RF-05, RF-07, RF-08'],
    ['CU-04', 'Consultar plan y consumo', 'Administrador de la empresa', W, '—', 'RF-06, RF-50'],
    ['CU-05', 'Verificar límites y funciones del plan', 'Sistema', BE, '—', 'RF-09'],
    ['CU-06', 'Iniciar sesión en una empresa', 'Todos los usuarios', WM, 'CU-02', 'RF-03'],
    ['CU-07', 'Registrar cliente con cuenta única', 'Cliente', WM, 'CU-01', 'RF-51'],
    ['CU-08', 'Gestionar usuarios de la empresa', 'Administrador de la empresa', W, 'CU-03', 'RF-03, RF-09'],
    ['CU-09', 'Gestionar sucursales de la empresa', 'Administrador de la empresa', W, 'CU-05', 'RF-03, RF-09'],
    ['CU-10', 'Gestionar empleados de la empresa', 'Administrador de la empresa', W, 'CU-06', 'RF-09'],
    ['CU-11', 'Contratar un plan desde la web (propio)', 'Empresa interesada', W, '—', 'RF-47'],
    ['CU-12', 'Personalizar la tienda de la empresa (propio)', 'Administrador de la empresa', W, '—', 'RF-46'],
    ['CU-13', 'Preparar el dispositivo para trabajar offline', 'Cajero', WM, '—', 'RF-10 a RF-13'],
    ['CU-14', 'Registrar venta offline', 'Cajero', WM, 'CU-31', 'RF-15'],
    ['CU-15', 'Registrar reserva offline', 'Cajero', WM, 'CU-22', 'RF-14'],
    ['CU-16', 'Sincronizar operaciones pendientes', 'Sistema', WM, '—', 'RF-18 a RF-21'],
    ['CU-17', 'Consultar operaciones pendientes', 'Cajero', WM, '—', 'RF-16, RF-17, RF-20, RF-23'],
    ['CU-18', 'Resolver conflictos de sincronización', 'Encargado de sucursal', W, '—', 'RF-22'],
    ['CU-19', 'Abrir y cerrar caja con ventas offline', 'Cajero', W, 'CU-30', 'RF-15, RF-21'],
    ['CU-20', 'Registrar cambio o devolución en caja', 'Cajero', W, 'CU-32', 'RF-24, RF-37'],
    ['CU-21', 'Configurar la ventana de devolución', 'Administrador de la empresa', W, '—', 'RF-35'],
    ['CU-22', 'Solicitar cambio o devolución por la web', 'Cliente', W, '—', 'RF-34'],
    ['CU-23', 'Atender solicitud de devolución', 'Encargado de sucursal', W, '—', 'RF-36, RF-37'],
    ['CU-24', 'Notificar eventos por correo según el plan', 'Sistema', BE, 'CU-40', 'RF-49'],
    ['CU-25', 'Generar reporte con IA por texto o voz', 'Administrador de la empresa', W, 'CU-35', 'RF-25 a RF-30'],
    ['CU-26', 'Generar reportes de gestión con devoluciones', 'Administrador de la empresa', W, 'CU-37', 'RF-38'],
    ['CU-27', 'Entrenar los modelos de analítica', 'Sistema', BE, '—', 'RF-31'],
    ['CU-28', 'Consultar el dashboard analítico', 'Administrador de la empresa', W, '—', 'RF-32, RF-33'],
    ['CU-29', 'Solicitar traspaso entre sucursales', 'Encargado de sucursal', W, 'CU-15', 'RF-39, RF-40'],
    ['CU-30', 'Atender traspaso', 'Encargado de sucursal', W, '—', 'RF-41'],
    ['CU-31', 'Consultar traspasos sugeridos (propio)', 'Administrador de la empresa', W, '—', 'RF-42'],
    ['CU-32', 'Preparar prenda para el probador (propio)', 'Administrador de la empresa', W, 'CU-11', 'RF-43'],
    ['CU-33', 'Probarse una prenda con su talla real (propio)', 'Cliente', M, 'CU-21', 'RF-44'],
    ['CU-34', 'Generar foto realista de la prueba (propio)', 'Cliente', M, '—', 'RF-45'],
    ['CU-35', 'Exportar los datos de la empresa (propio)', 'Administrador de la empresa', W, '—', 'RF-48'],
]
h25 = parrafo_que_empieza('2.5. Diagrama UML de casos de uso')
intro = h25.getnext()
poner_texto(intro, 'Son 35 casos de uso, numerados desde el CU-01. Los del primer parcial que no se rehacen se '
                   'listan en el Anexo 1, junto con su diagrama, y se citan como «CU-NN del primer parcial». '
                   'Los que tienen algo en la columna Reemplaza vuelven a especificarse y a modelarse completos, '
                   'porque el multitenant, los planes o el modo offline cambian su flujo. Los marcados como '
                   'propios son aporte del grupo. La columna Plataforma indica si el caso de uso se usa en la '
                   'web, en el móvil o en ambos; los del backend no tienen pantalla propia.')
segundo = intro.getnext()  # «Todos los RF ... incluyen al CU-47»: se conserva y se renumera más abajo
fin = parrafo_que_empieza('[Pendiente] Diagrama de casos de uso por paquete', segundo)
el = segundo.getnext()
while el is not fin:
    sig = el.getnext()
    body.remove(el)
    el = sig
lista = tabla(['CU', 'Caso de uso', 'Actor principal', 'Plataforma', 'Reemplaza (1er parcial)', 'RF'],
              CASOS, [1150, 2632, 1900, 1200, 1500, 1400])
insertar_despues(segundo, [lista, nuevo_parrafo()])

# =====================================================================
# 4) Renumerar CU-43..CU-77 -> CU-01..CU-35 en todo el cuerpo (los del 1er parcial llegan hasta el CU-42)
def renumerar(m):
    n = int(m.group(1))
    return f'CU-{n - 42:02d}' if 43 <= n <= 77 else m.group(0)


cambios = 0
for el in list(body.iter(qn('w:p'))):
    if el.getparent() is not None and el.getparent().tag == qn('w:tc') and el in lista.iter(qn('w:p')):
        continue  # la lista nueva ya viene numerada
    cambios += reemplazar_en_parrafo(el, r'\bCU-(\d\d)\b', renumerar)
print('referencias renumeradas', cambios)

# =====================================================================
# 5) Otras menciones a los ciclos
for el in body.iter(qn('w:p')):
    reemplazar_en_parrafo(el, r'ICONIX en tres ciclos',
                          lambda m: 'ICONIX, con casos de uso que se agregan en cada presentación')
    reemplazar_en_parrafo(el, r'Se escriben al cerrar el Ciclo 3\.',
                          lambda m: 'Se escriben para la presentación final.')

# =====================================================================
# 6) Anexo 1: lista y diagramas de casos de uso del primer parcial, copiados del documento de la defensa
p1 = docx.Document(P1)
b1 = p1.element.body
h112 = next(e for e in b1.iterchildren() if e.tag == qn('w:p') and texto(e).strip() == '1.1.2 Casos de Uso')
lista_p1 = h112.getnext()
while lista_p1.tag != qn('w:tbl'):
    lista_p1 = lista_p1.getnext()
h15 = next(e for e in b1.iterchildren() if e.tag == qn('w:p') and texto(e).strip().startswith('1.5 Estructurar'))
imagenes = []
el = h15.getnext()
while not (el.tag == qn('w:p') and texto(el).startswith('CAP. 2')):
    if el.tag == qn('w:p') and list(el.iter(qn('w:drawing'))):
        imagenes.append(el)
    el = el.getnext()
assert len(imagenes) == 3, len(imagenes)

_doc_pr = [9000]


def copiar_imagen(par):
    nuevo = copy.deepcopy(par)
    for blip in nuevo.iter(qn('a:blip')):
        rid = blip.get(qn('r:embed'))
        blob = p1.part.related_parts[rid].blob
        nuevo_rid, _ = doc.part.get_or_add_image(io.BytesIO(blob))
        blip.set(qn('r:embed'), nuevo_rid)
    for dp in nuevo.iter(qn('wp:docPr')):
        _doc_pr[0] += 1
        dp.set('id', str(_doc_pr[0]))
    return nuevo


def poner_celda(tc, valor):
    ps = tc.findall(qn('w:p'))
    for p in ps[1:]:
        tc.remove(p)
    ts = list(ps[0].iter(qn('w:t')))
    ts[0].text = valor
    for t in ts[1:]:
        t.text = ''


# la tabla de la lista en el documento de la defensa no trae la fila del CU-42; se agrega
lista_p1 = copy.deepcopy(lista_p1)
filas = lista_p1.findall(qn('w:tr'))
if texto(filas[-1]).startswith('CU-41'):
    fila42 = copy.deepcopy(filas[-2])  # la penúltima, para seguir la cebra
    for tc, valor in zip(fila42.findall(qn('w:tc')),
                         ['CU-42', 'Consultar la bitácora del sistema', 'Administrador']):
        poner_celda(tc, valor)
    filas[-1].addnext(fila42)

anexos = parrafo_que_empieza('8. ANEXOS')
pend_enlaces = parrafo_que_empieza('[Pendiente] Enlaces al sistema desplegado', anexos)
piezas = [
    nuevo_parrafo('Anexo 1. Casos de uso del primer parcial', 'Titulo1'),
    nuevo_parrafo('Los 42 casos de uso del primer parcial siguen funcionando en el sistema. Los que el '
                  'multitenant, los planes o el modo offline modifican se vuelven a especificar como casos de uso '
                  'nuevos (columna Reemplaza de la sección 2.5); el resto se conserva tal como se defendió y se '
                  'cita como «CU-NN del primer parcial».'),
    nuevo_parrafo('Lista de casos de uso', 'Titulo3'),
    lista_p1,
    nuevo_parrafo(),
    nuevo_parrafo('Diagrama de casos de uso', 'Titulo3'),
]
Paragraph(piezas[-1], None).paragraph_format.keep_with_next = True
for k, img in enumerate(imagenes, 1):
    pie = nuevo_parrafo(f'Modelo de casos de uso estructurado del primer parcial, parte {k} de 3.')
    Paragraph(pie, None).paragraph_format.keep_with_next = True  # que no quede solo al pie de la página
    piezas.append(pie)
    piezas.append(copiar_imagen(img))
piezas += [nuevo_parrafo('Anexo 2. Enlaces', 'Titulo1')]
insertar_despues(anexos, piezas)
# el pendiente de enlaces queda bajo el Anexo 2
piezas[-1].addnext(pend_enlaces)

# el sectPr tiene que quedar último
sect = body.find(qn('w:sectPr'))
body.remove(sect); body.append(sect)
doc.save(DOC)
print('ok', DOC)
