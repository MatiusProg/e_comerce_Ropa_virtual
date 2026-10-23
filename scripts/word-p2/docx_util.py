# Herramientas para editar en el lugar el Word del 2do parcial, que ya tiene trabajo hecho a mano
# (por eso no se regenera con gen_doc.py). Las tablas salen con el mismo formato que gen_doc.py.
import copy, re
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

BORDE = 'D1D9E0'
ROJO = RGBColor(0xEE, 0x00, 0x00)
XML_SPACE = '{http://www.w3.org/XML/1998/namespace}space'


class Editor:
    def __init__(self, doc):
        self.doc = doc
        self.body = doc.element.body

    # ---------- búsqueda ----------
    @staticmethod
    def texto(el):
        return ''.join(t.text or '' for t in el.iter(qn('w:t')))

    def buscar(self, pred, desde=None):
        hijos = list(self.body.iterchildren())
        i0 = hijos.index(desde) + 1 if desde is not None else 0
        for el in hijos[i0:]:
            if pred(el):
                return el
        raise LookupError('no encontrado')

    def parrafo(self, prefijo, desde=None, exacto=False):
        def pred(e):
            if e.tag != qn('w:p'):
                return False
            t = self.texto(e).strip()
            return t == prefijo if exacto else t.startswith(prefijo)
        return self.buscar(pred, desde)

    def tabla_despues(self, el):
        while el.tag != qn('w:tbl'):
            el = el.getnext()
        return el

    # ---------- texto ----------
    @staticmethod
    def reemplazar(p, patron, funcion):
        """Reemplaza sobre el texto completo del párrafo aunque esté partido en varios runs."""
        ts = list(p.iter(qn('w:t')))
        if not ts:
            return 0
        completo = ''.join(t.text or '' for t in ts)
        coincidencias = list(re.finditer(patron, completo))
        if not coincidencias:
            return 0
        mapa = [(k, j) for k, t in enumerate(ts) for j in range(len(t.text or ''))]
        textos = [list(t.text or '') for t in ts]
        for m in reversed(coincidencias):
            nuevo = funcion(m) if callable(funcion) else funcion
            for pos in range(m.end() - 1, m.start(), -1):
                k, j = mapa[pos]
                textos[k][j] = ''
            k, j = mapa[m.start()]
            textos[k][j] = nuevo
        for t, chars in zip(ts, textos):
            t.text = ''.join(chars)
            t.set(XML_SPACE, 'preserve')
        return len(coincidencias)

    def reemplazar_en_todo(self, patron, funcion):
        return sum(self.reemplazar(p, patron, funcion) for p in list(self.body.iter(qn('w:p'))))

    @staticmethod
    def poner_texto(p, nuevo):
        """Deja el párrafo con un único texto (**negrita** marcada), con el formato del primer run."""
        runs = p.findall(qn('w:r'))
        base = copy.deepcopy(runs[0]) if runs else None
        for r in runs:
            p.remove(r)
        for i, parte in enumerate(nuevo.split('**')):
            if not parte:
                continue
            r = copy.deepcopy(base) if base is not None else OxmlElement('w:r')
            for hijo in list(r):
                if hijo.tag != qn('w:rPr'):
                    r.remove(hijo)
            rpr = r.find(qn('w:rPr'))
            if rpr is None:
                rpr = OxmlElement('w:rPr'); r.insert(0, rpr)
            for b in rpr.findall(qn('w:b')):
                rpr.remove(b)
            if i % 2 == 1:
                rpr.insert(0, OxmlElement('w:b'))
            t = OxmlElement('w:t'); t.text = parte; t.set(XML_SPACE, 'preserve')
            r.append(t)
            p.append(r)

    def poner_celda(self, tc, valor):
        """Reemplaza el texto de una celda conservando el formato de su primer párrafo y su primer run."""
        ps = tc.findall(qn('w:p'))
        for p in ps[1:]:
            tc.remove(p)
        self.poner_texto(ps[0], valor)

    # ---------- piezas nuevas (se crean al final y se mueven) ----------
    @staticmethod
    def _runs(p, txt, tam=None):
        for i, parte in enumerate(txt.split('**')):
            if parte:
                r = p.add_run(parte)
                r.bold = (i % 2 == 1) or None
                if tam:
                    r.font.size = Pt(tam)

    def nuevo_parrafo(self, txt='', estilo=None, con_siguiente=False):
        p = self.doc.add_paragraph(style=estilo) if estilo else self.doc.add_paragraph()
        if txt:
            if estilo:
                p.add_run(txt)
            else:
                p.paragraph_format.line_spacing = 1.15
                p.alignment = 3  # justificado
                self._runs(p, txt)
        if con_siguiente:
            p.paragraph_format.keep_with_next = True
        return p._p

    def copiar_parrafo(self, modelo, txt):
        """Copia un párrafo existente (estilo y formato) con otro texto; sirve para títulos hechos a mano."""
        nuevo = copy.deepcopy(modelo)
        self.poner_texto(nuevo, txt)
        return nuevo

    def _celda(self, celda, txt, ancho, encabezado, cebra, negrita):
        tcPr = celda._tc.get_or_add_tcPr()
        for hijo in list(tcPr):
            tcPr.remove(hijo)
        w = OxmlElement('w:tcW'); w.set(qn('w:w'), str(ancho)); w.set(qn('w:type'), 'dxa'); tcPr.append(w)
        b = OxmlElement('w:tcBorders')
        for lado in ('top', 'left', 'bottom', 'right'):
            e = OxmlElement(f'w:{lado}')
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '6'); e.set(qn('w:space'), '0')
            e.set(qn('w:color'), BORDE)
            b.append(e)
        tcPr.append(b)
        s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto')
        s.set(qn('w:fill'), 'F6F8FA' if cebra else 'FFFFFF'); tcPr.append(s)
        m = OxmlElement('w:tcMar')
        for lado, v in (('top', 90), ('left', 195), ('bottom', 90), ('right', 195)):
            e = OxmlElement(f'w:{lado}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
        tcPr.append(m)
        va = OxmlElement('w:vAlign'); va.set(qn('w:val'), 'center'); tcPr.append(va)
        p = celda.paragraphs[0]
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.first_line_indent = 0
        if encabezado:
            r = p.add_run(txt); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = ROJO
        else:
            self._runs(p, f'**{txt}**' if negrita and '**' not in txt else txt, tam=11)

    def tabla(self, encabezados, filas, anchos):
        total = 9782
        assert sum(anchos) == total, (encabezados, sum(anchos))
        t = self.doc.add_table(rows=1 + len(filas), cols=len(encabezados))
        tblPr = t._tbl.tblPr
        for tag in ('w:tblStyle', 'w:tblW', 'w:tblInd', 'w:tblLook'):
            for e in tblPr.findall(qn(tag)):
                tblPr.remove(e)
        tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(total)); tw.set(qn('w:type'), 'dxa'); tblPr.append(tw)
        ti = OxmlElement('w:tblInd'); ti.set(qn('w:w'), '-434'); ti.set(qn('w:type'), 'dxa'); tblPr.append(ti)
        for gc, a in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), anchos):
            gc.set(qn('w:w'), str(a))
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for j, h in enumerate(encabezados):
            self._celda(t.rows[0].cells[j], h, anchos[j], True, False, False)
        for i, fila in enumerate(filas, start=1):
            for j, v in enumerate(fila):
                self._celda(t.rows[i].cells[j], v, anchos[j], False, i % 2 == 0, j == 0)
        return t._tbl

    # ---------- filas de tablas existentes ----------
    def agregar_fila(self, tbl, valores, despues_de=None):
        """Agrega una fila copiando el formato de una existente; luego hay que llamar a recebrar()."""
        filas = tbl.findall(qn('w:tr'))
        ancla = despues_de if despues_de is not None else filas[-1]
        nueva = copy.deepcopy(filas[-1])
        for j, (tc, v) in enumerate(zip(nueva.findall(qn('w:tc')), valores)):
            self.poner_celda(tc, f'**{v}**' if j == 0 else v)  # la primera columna va en negrita
        ancla.addnext(nueva)
        return nueva

    @staticmethod
    def recebrar(tbl):
        for i, tr in enumerate(tbl.findall(qn('w:tr'))):
            if i == 0:
                continue
            for shd in tr.iter(qn('w:shd')):
                shd.set(qn('w:fill'), 'F6F8FA' if i % 2 == 0 else 'FFFFFF')

    def fila(self, tbl, primera_celda):
        for tr in tbl.findall(qn('w:tr')):
            if self.texto(tr.find(qn('w:tc'))).strip() == primera_celda:
                return tr
        raise LookupError(primera_celda)

    @staticmethod
    def insertar_despues(ancla, elementos):
        for el in elementos:
            ancla.addnext(el)
            ancla = el
        return ancla

    def cerrar(self, ruta):
        sect = self.body.find(qn('w:sectPr'))
        self.body.remove(sect); self.body.append(sect)
        self.doc.save(ruta)
