# Paso 3 de las correcciones del 22/10, sobre el Word ya editado a mano:
#   - CU nuevos pedidos por la ing: inventario, pago sin conexión, ofertas y precios mayoristas;
#   - el cambio de prenda considera que cada talla puede tener otro precio;
#   - regla de Karen para el offline: gana el primero que llega al servidor y al que pagó sin conexión se le
#     devuelve el dinero;
#   - los CU nuevos entran en el bloque de su presentación (se renumeran solo los CU-13 en adelante, que todavía
#     no tienen fichas ni diagramas);
#   - el CU-12 que faltaba en los subtítulos del capítulo III y las descripciones del CU-38 al CU-42 del Anexo 1.
# Uso: python editar_23_10.py <documento.docx> <documento_defensa_1er_parcial.docx>
import copy, sys
import docx
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx_util import Editor

DOC, P1 = sys.argv[1], sys.argv[2]
ed = Editor(docx.Document(DOC))
T = ed.texto

# =====================================================================
# Lista nueva de casos de uso (por presentación, en bloques contiguos)
W, M, WM, BE = 'Web', 'Móvil', 'Web y móvil', 'Backend'
CASOS = [
    # Presentación 1 (sin cambios)
    ['CU-01', 'Gestionar empresas', 'Administrador de la plataforma', W, '—', 'RF-01, RF-02, RF-03'],
    ['CU-02', 'Gestionar planes', 'Administrador de la plataforma', W, '—', 'RF-04'],
    ['CU-03', 'Gestionar suscripciones', 'Administrador de la plataforma', W, '—', 'RF-05, RF-07, RF-08'],
    ['CU-04', 'Consultar plan y consumo', 'Administrador de la empresa', W, '—', 'RF-06, RF-50'],
    ['CU-05', 'Verificar límites y funciones del plan', 'Sistema', BE, '—', 'RF-09'],
    ['CU-06', 'Iniciar sesión en una empresa', 'Todos los usuarios', WM, 'CU-02', 'RF-03'],
    ['CU-07', 'Registrar cliente con cuenta única', 'Cliente', WM, 'CU-01', 'RF-51'],
    ['CU-08', 'Gestionar usuarios de la empresa', 'Administrador de la empresa', W, 'CU-03', 'RF-03, RF-09'],
    ['CU-09', 'Gestionar sucursales de la empresa', 'Administrador de la empresa', W, 'CU-05', 'RF-03, RF-09'],
    ['CU-10', 'Gestionar empleados de la empresa', 'Administrador de la empresa', W, 'CU-06', 'RF-09'],
    ['CU-11', 'Contratar un plan desde la web (propio)', 'Empresa interesada', W, '—', 'RF-47'],
    ['CU-12', 'Personalizar la tienda de la empresa (propio)', 'Administrador de la empresa', W, '—', 'RF-46'],
    # Presentación 2
    ['CU-13', 'Preparar el dispositivo para trabajar offline', 'Cajero', WM, '—', 'RF-10 a RF-13'],
    ['CU-14', 'Registrar venta offline', 'Cajero', WM, 'CU-31', 'RF-15'],
    ['CU-15', 'Registrar reserva offline', 'Cajero o cliente', WM, 'CU-22', 'RF-14'],
    ['CU-16', 'Registrar pago sin conexión', 'Cajero o cliente', WM, '—', 'RF-54'],
    ['CU-17', 'Sincronizar operaciones pendientes', 'Sistema', WM, '—', 'RF-18 a RF-21'],
    ['CU-18', 'Consultar operaciones pendientes', 'Cajero', WM, '—', 'RF-16, RF-17, RF-20, RF-23'],
    ['CU-19', 'Atender operaciones rechazadas y devolver el dinero', 'Encargado de sucursal', W, '—',
     'RF-22, RF-55'],
    ['CU-20', 'Abrir y cerrar caja con ventas offline', 'Cajero', W, 'CU-30', 'RF-15, RF-21'],
    ['CU-21', 'Gestionar el inventario de la sucursal', 'Encargado de sucursal', WM, 'CU-13, CU-15, CU-16',
     'RF-52, RF-53'],
    ['CU-22', 'Registrar cambio o devolución en caja', 'Cajero', W, 'CU-32', 'RF-24, RF-37, RF-59'],
    ['CU-23', 'Configurar la ventana de devolución', 'Administrador de la empresa', W, '—', 'RF-35'],
    ['CU-24', 'Solicitar cambio o devolución por la web', 'Cliente', W, '—', 'RF-34, RF-59'],
    ['CU-25', 'Atender solicitud de devolución', 'Encargado de sucursal', W, '—', 'RF-36, RF-37'],
    ['CU-26', 'Notificar eventos por correo según el plan', 'Sistema', BE, 'CU-40', 'RF-49'],
    # Presentación 3
    ['CU-27', 'Generar reporte con IA por texto o voz', 'Administrador de la empresa', W, 'CU-35', 'RF-25 a RF-30'],
    ['CU-28', 'Generar reportes de gestión con devoluciones', 'Administrador de la empresa', W, 'CU-37', 'RF-38'],
    ['CU-29', 'Entrenar los modelos de analítica', 'Sistema', BE, '—', 'RF-31'],
    ['CU-30', 'Consultar el dashboard analítico', 'Administrador de la empresa', W, '—', 'RF-32, RF-33'],
    ['CU-31', 'Gestionar ofertas y precios mayoristas', 'Administrador de la empresa', W, 'CU-12',
     'RF-56, RF-58'],
    ['CU-32', 'Aplicar ofertas de fuera de temporada', 'Sistema', BE, '—', 'RF-57'],
    ['CU-33', 'Solicitar traspaso entre sucursales', 'Encargado de sucursal', W, 'CU-15', 'RF-39, RF-40'],
    ['CU-34', 'Atender traspaso', 'Encargado de sucursal', W, '—', 'RF-41'],
    ['CU-35', 'Consultar traspasos sugeridos (propio)', 'Administrador de la empresa', W, '—', 'RF-42'],
    ['CU-36', 'Preparar prenda para el probador (propio)', 'Administrador de la empresa', W, 'CU-11', 'RF-43'],
    ['CU-37', 'Probarse una prenda con su talla real (propio)', 'Cliente', M, 'CU-21', 'RF-44'],
    ['CU-38', 'Generar foto realista de la prueba (propio)', 'Cliente', M, '—', 'RF-45'],
    ['CU-39', 'Exportar los datos de la empresa (propio)', 'Administrador de la empresa', W, '—', 'RF-48'],
]
NOMBRE = {c[0]: c[1].replace(' (propio)', '') for c in CASOS}  # los títulos hechos a mano no lo llevan

# =====================================================================
# 1.3.2, 1.5: objetivos, alcance y limitaciones
p = ed.parrafo('Implementar las devoluciones vía web pedidas en clase')
ed.poner_texto(p, 'Implementar lo pedido por la docente en clase (devoluciones vía web, inventario por sucursal, '
                  'pago sin conexión, ofertas con precios mayoristas y cambios de prenda con diferencia de '
                  'precio) y los aportes propios del grupo: traspasos entre sucursales y un probador virtual que '
                  'distinga tallas.')
p = ed.parrafo('Módulo 11 — Aportes propios.')
ed.insertar_despues(p, [ed.nuevo_parrafo(
    '**Módulo 12 — Inventario, pago sin conexión y ofertas.** Pedido de la docente en clase: inventario por '
    'sucursal con su kardex, también sin conexión; pago de ventas y reservas sin conexión, con la devolución del '
    'dinero si al sincronizar la prenda ya no está; precios mayoristas desde una cantidad mínima, ofertas '
    'manuales y ofertas automáticas de fuera de temporada; y cambios de prenda que cobran o devuelven la '
    'diferencia de precio entre tallas.')])
p = ed.parrafo('Módulo 7 — Ventas offline y cambios de prenda.')
ed.poner_texto(p, '**Módulo 7 — Ventas offline y cambios de prenda.** El servidor confirma o rechaza cada venta '
                  'y cada pago offline contra el inventario central, y se gestionan los cambios de prenda dentro '
                  'de un plazo.')
p = ed.parrafo('Las ventas offline son solo de contado')
ed.poner_texto(p, 'Sin conexión se cobra en efectivo o con QR, anotando la referencia del pago; la pasarela de '
                  'pago en línea sigue necesitando conexión.')

# =====================================================================
# 1.6: rangos de cada presentación
tbl = ed.tabla_despues(ed.parrafo('Cómo se aplica en este proyecto.'))
for nombre, texto in (
        ('Presentación 2', 'CU-13 a CU-26: contrato de sincronización; PWA con IndexedDB; Flutter con SQLite; '
                           'ventas, reservas y pagos offline; inventario de la sucursal; cambios de prenda, '
                           'devoluciones web y avisos por correo'),
        ('Presentación 3', 'CU-27 a CU-39: reportes con IA y voz; analítica con ML; ofertas y precios '
                           'mayoristas; traspasos y probador por talla; exportación de datos')):
    ed.poner_celda(ed.fila(tbl, nombre).findall(qn('w:tc'))[2], texto)

# =====================================================================
# 2.1.1: planes (lo pedido en clase va en todos)
tbl = ed.tabla_despues(ed.parrafo('2.1.1. Planes de suscripción'))
ancla = ed.fila(tbl, 'Traspasos entre sucursales a pedido')
for fila in (['Inventario por sucursal con kardex, también sin conexión', 'Sí', 'Sí', 'Sí'],
             ['Pago sin conexión (efectivo o QR)', 'Sí', 'Sí', 'Sí'],
             ['Ofertas manuales y precios mayoristas', 'Sí', 'Sí', 'Sí'],
             ['Ofertas automáticas de fuera de temporada', 'Sí', 'Sí', 'Sí']):
    ancla = ed.agregar_fila(tbl, fila, despues_de=ancla)
ed.recebrar(tbl)

# =====================================================================
# 2.3.2: requisitos pedidos en clase
p = ed.parrafo('La docente los pidió en clase')
ed.poner_texto(p, 'La docente los pidió en clase, aunque no figuran en el enunciado. En caja el cambio de prenda '
                  'ya funciona desde el primer parcial (CU-32 del primer parcial); ahora el cliente también puede '
                  'pedirlo desde la tienda en línea. También pidió un caso de uso de inventario, el pago sin '
                  'conexión, las ofertas con precios mayoristas y que el cambio de prenda considere que una talla '
                  'puede costar distinto que otra.')
tbl_dev = ed.tabla_despues(ed.parrafo('Devoluciones vía web'))
A2 = [1100, 8682]
nuevas = [
    ed.nuevo_parrafo('Inventario por sucursal', 'Titulo3', con_siguiente=True),
    ed.tabla(['RF', 'El sistema deberá permitir…'], [
        ['RF-52', 'Registrar ingresos, salidas y ajustes de inventario por sucursal, con su motivo, su usuario y '
                  'su fecha, también sin conexión (se sincronizan después)'],
        ['RF-53', 'Consultar la existencia por sucursal y el kardex de cada prenda: todos sus movimientos, '
                  'incluidas ventas, reservas, traspasos y notas de devolución'],
    ], A2),
    ed.nuevo_parrafo(),
    ed.nuevo_parrafo('Pago sin conexión', 'Titulo3', con_siguiente=True),
    ed.tabla(['RF', 'El sistema deberá permitir…'], [
        ['RF-54', 'Registrar el pago de una venta o una reserva sin conexión, en efectivo o con QR anotando la '
                  'referencia, para que el servidor lo confirme al sincronizar'],
        ['RF-55', 'Si al sincronizar la prenda ya fue vendida o reservada por alguien con conexión, rechazar la '
                  'operación offline y registrar la devolución del dinero al cliente'],
    ], A2),
    ed.nuevo_parrafo(),
    ed.nuevo_parrafo('Ofertas y precios mayoristas', 'Titulo3', con_siguiente=True),
    ed.tabla(['RF', 'El sistema deberá permitir…'], [
        ['RF-56', 'Definir precios mayoristas que se aplican desde una cantidad mínima de prendas, y mostrarlos en '
                  'el catálogo y en la caja'],
        ['RF-57', 'Aplicar automáticamente ofertas de fuera de temporada según un factor medible guardado en la '
                  'base: días desde el fin de la temporada y existencia sobrante'],
        ['RF-58', 'Seguir creando ofertas manuales, con su vigencia, sobre una prenda, una categoría o una '
                  'temporada'],
    ], A2),
    ed.nuevo_parrafo(),
    ed.nuevo_parrafo('Cambio de prenda con diferencia de precio', 'Titulo3', con_siguiente=True),
    ed.tabla(['RF', 'El sistema deberá permitir…'], [
        ['RF-59', 'En un cambio de prenda, comparar el precio pagado con el de la prenda o talla nueva, y cobrar '
                  'la diferencia o devolverla'],
    ], A2),
    ed.nuevo_parrafo(),
]
fin_dev = tbl_dev.getnext()  # el párrafo vacío que sigue a la tabla
ed.insertar_despues(fin_dev, nuevas)

# =====================================================================
# 2.4: reglas de negocio
tbl = ed.tabla_despues(ed.parrafo('2.4. Reglas de negocio'))
ed.poner_celda(ed.fila(tbl, 'RN-05').findall(qn('w:tc'))[1],
               'Una venta o una reserva sin conexión se paga en efectivo o con QR, anotando la referencia. Cobrar '
               'no la confirma: la confirma el servidor al sincronizar')
ed.poner_celda(ed.fila(tbl, 'RN-06').findall(qn('w:tc'))[1],
               'Si al sincronizar ya no alcanza la existencia, gana la operación que llegó primero al servidor; la '
               'que se hizo sin conexión queda rechazada y, si ya se cobró, se le devuelve el dinero al cliente')
for fila in (
        ['RN-15', 'Sobre una prenda se aplica un solo precio: el más bajo entre el normal, el mayorista (si se '
                  'alcanza la cantidad mínima) y la oferta vigente, manual o automática; los descuentos no se '
                  'acumulan'],
        ['RN-16', 'Una oferta automática de fuera de temporada se activa solo cuando se cumple el factor '
                  'configurado (días desde el fin de la temporada y existencia sobrante); el administrador puede '
                  'desactivarla'],
        ['RN-17', 'En un cambio de prenda, si la prenda o talla nueva cuesta más se cobra la diferencia, y si '
                  'cuesta menos se devuelve']):
    ed.agregar_fila(tbl, fila)
ed.recebrar(tbl)

# =====================================================================
# 2.5: lista de casos de uso
h25 = ed.parrafo('2.5. Diagrama UML de casos de uso')
intro = h25.getnext()
ed.reemplazar(intro, r'Son 35 casos de uso', 'Son 39 casos de uso')
segundo = intro.getnext()
ed.reemplazar(segundo, r'del RF-01 al RF-51', 'del RF-01 al RF-59')
vieja = ed.tabla_despues(segundo)
nueva = ed.tabla(['CU', 'Caso de uso', 'Actor principal', 'Plataforma', 'Reemplaza (1er parcial)', 'RF'],
                 CASOS, [1150, 2632, 1900, 1200, 1500, 1400])
vieja.addprevious(nueva)
ed.body.remove(vieja)

# =====================================================================
# 2.6: los títulos de los CU de las presentaciones 2 y 3, con la numeración nueva
h_sig = ed.parrafo('PARA SIGUIENTES PRESENTACIONES', exacto=True)
modelo = copy.deepcopy(h_sig.getnext())  # un título «CU-13» hecho a mano
cap3 = ed.parrafo('CAPÍTULO III')
separador = None  # el renglón vacío que separaba las presentaciones 2 y 3
el = h_sig.getnext()
while el is not cap3:
    sig = el.getnext()
    if separador is None and el.tag == qn('w:p') and not T(el).strip() and not list(el.iter(qn('w:br'))):
        separador = copy.deepcopy(el)
    ed.body.remove(el)  # el último título traía el salto de página: se repone abajo
    el = sig
titulos = [ed.copiar_parrafo(modelo, f'{c} — {NOMBRE[c]}') for c, *_ in CASOS[12:]]
if separador is not None:
    titulos.insert(26 - 12, separador)  # antes del CU-27
salto = ed.doc.add_paragraph()
salto.add_run().add_break(WD_BREAK.PAGE)
ed.insertar_despues(h_sig, titulos + [salto._p])

# =====================================================================
# Capítulo III: el CU-12 que faltaba después del CU-11 en cada sección
cap4 = ed.parrafo('CAPÍTULO IV')
hijos = list(ed.body.iterchildren())
agregados = 0
for el in hijos[hijos.index(cap3):hijos.index(cap4)]:
    if el.tag == qn('w:p') and T(el).strip().startswith('CU-11 —'):
        sig = el.getnext()
        if not T(sig).strip().startswith('CU-12'):
            el.addnext(ed.copiar_parrafo(el, f'CU-12 — {NOMBRE["CU-12"]}'))
            agregados += 1
print('títulos CU-12 agregados', agregados)

# =====================================================================
# Anexo 1: las filas del CU-38 al CU-42 traían el actor en lugar de la descripción
p1 = docx.Document(P1)
desc = {}
for t in p1.tables:
    try:
        c = t.rows[0].cells
    except IndexError:
        continue
    if len(c) > 1:
        for n in range(38, 43):
            if c[1].text.startswith(f'CU-{n} '):
                for r in t.rows:
                    if r.cells[0].text.strip() == 'Descripción':
                        desc[f'CU-{n}'] = r.cells[1].text.strip()
anexo = ed.tabla_despues(ed.parrafo('Anexo 1. Casos de uso del primer parcial'))
for cu, d in desc.items():
    ed.poner_celda(ed.fila(anexo, cu).findall(qn('w:tc'))[2], d)
print('descripciones del anexo', sorted(desc))

ed.cerrar(DOC)
print('ok', DOC)
